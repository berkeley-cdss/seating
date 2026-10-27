import csv
import io
import zipfile

from openpyxl import load_workbook

from server.typings.exception import DataValidationError


class LowerCaseDictReader(csv.DictReader):
    """
    Custom CSV DictReader that converts all headers to lowercase first.
    """

    def __init__(self, *args, **kwargs):
        super(LowerCaseDictReader, self).__init__(*args, **kwargs)
        self.fieldnames = [field.lower() for field in self.fieldnames]


def parse_csv(file):
    """
    Parse a CSV file and return a list of headers and rows.
    """
    # utf-8-sig also reads Excel's "CSV UTF-8" files, which start with a byte order mark
    return parse_csv_str(file.read().decode('utf-8-sig'))


def parse_csv_str(csv_str, delimiter=','):
    """
    Parse a CSV string and return a list of headers and rows.
    """
    reader = LowerCaseDictReader(io.StringIO(csv_str), delimiter=delimiter)
    headers = list(reader.fieldnames)
    rows = [row for row in reader]
    return headers, rows


def parse_xlsx(content):
    """
    Parse the first sheet of an Excel (.xlsx) file, like parse_csv_str: lowercase headers, and every
    cell a string. Excel stores numbers as floats, so whole numbers (canvas ids, say) lose their ".0",
    and TRUE and FALSE cells read as "true" and "false". Empty lines are skipped.
    """
    try:
        sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).worksheets[0]
    except (zipfile.BadZipFile, KeyError, IndexError, OSError):
        raise DataValidationError("Couldn't read the Excel file: save it as .xlsx or as CSV")

    def cell(value):
        if value is None:
            return ''
        if isinstance(value, bool):
            return str(value).lower()
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value).strip()

    lines = [[cell(value) for value in line] for line in sheet.iter_rows(values_only=True)]
    lines = [line for line in lines if any(line)]
    if not lines:
        return [], []
    headers = [header.lower() for header in lines[0]]
    rows = [{header: value for header, value in zip(headers, line) if header} for line in lines[1:]]
    return [header for header in headers if header], rows


def original_csv_headers(csv_str, delimiter=','):
    """
    The header row exactly as written, before parse_csv_str lowercases it.

    Useful when a column name is data - a layout is named by whoever made the
    spreadsheet, so it should be shown the way they spelled it.
    """
    return next(csv.reader(io.StringIO(csv_str), delimiter=delimiter), [])


def sniff_delimiter(text, candidates='\t,;'):
    """
    Guess what separates the columns of a delimited text blob.

    A range copied out of Google Sheets is tab separated, and a file saved from
    the same sheet is comma separated, so anything that takes both has to work
    it out. csv.Sniffer is unreliable here because most cells in a room layout
    are empty; the header row is the one line that is always fully populated, so
    count separators in that instead. Tabs win a tie: a heading like
    "2/3 Seating, Broken" contains a comma, but a heading never contains a tab.
    """
    header = text.lstrip('\ufeff').lstrip().split('\n', 1)[0]
    best = max(candidates, key=header.count)
    return best if header.count(best) else ','


def to_csv_str(headers, rows) -> str:
    """
    Convert a list of headers and rows to a CSV string.
    """
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=headers)
    writer.writeheader()
    for row in rows:
        writer.writerow(row)
    return output.getvalue()
